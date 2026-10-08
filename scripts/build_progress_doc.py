from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "DFI_Project_Progress_Evidence.docx"

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5B6573"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
GOLD = "7A5A00"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, name="Calibri", size=11, color="000000", bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_doc(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string("000000")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True


def add_header_footer(doc):
    sec = doc.sections[0]
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header.add_run("DFI PROJECT | IMPLEMENTATION PROGRESS EVIDENCE")
    set_run_font(header_run, size=9, color=MUTED, bold=True)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("Digital Forgetting Index  |  Page ")
    set_run_font(r, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)


def add_title(doc, text, subtitle=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_run_font(r, size=25, color=NAVY, bold=True)
    if subtitle:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(18)
        r2 = p2.add_run(subtitle)
        set_run_font(r2, size=14, color=MUTED)


def add_meta_table(doc):
    table = doc.add_table(rows=5, cols=2)
    set_table_geometry(table, [1800, 7560])
    rows = [
        ("Document type", "Project progress evidence and implementation record"),
        ("Project", "Digital Forgetting Index (DFI)"),
        ("Research basis", "Digital Forgetting Index: A Quantifiable Framework for Measuring AI Data Erasure"),
        ("Current status", "Reference implementation complete; real system results pending"),
        ("Prepared", "20 August 2026"),
    ]
    for i, (label, value) in enumerate(rows):
        c1, c2 = table.rows[i].cells
        set_cell_shading(c1, LIGHT_BLUE)
        c1.paragraphs[0].add_run(label)
        set_run_font(c1.paragraphs[0].runs[0], size=10, color=NAVY, bold=True)
        c2.paragraphs[0].add_run(value)
        set_run_font(c2.paragraphs[0].runs[0], size=10)
    doc.add_paragraph()


def add_callout(doc, label, text, fill=CALLOUT, color=NAVY):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label + "  ")
    set_run_font(r, size=10.5, color=color, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color="333333")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5 + level * 0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(5)
    p.add_run(text)
    return p


def add_table(doc, headers, rows, widths, header_fill=LIGHT_BLUE, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=font_size, color=NAVY, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color="222222")
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_result_space(doc, title, rows=4):
    doc.add_heading(title, level=2)
    p = doc.add_paragraph("Complete this section after running the model against the target system. Record the measured evidence, not only the final score.")
    p.paragraph_format.space_after = Pt(6)
    data = [("Measurement date", ""), ("System/model/database version", ""), ("Probe version", ""), ("Scope and exclusions", "")]
    data += [("Result", "") for _ in range(rows)]
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [2400, 6960])
    for cell, value in zip(table.rows[0].cells, ["Field", "Evidence / value to be added"]):
        set_cell_shading(cell, LIGHT_BLUE)
        r = cell.paragraphs[0].add_run(value)
        set_run_font(r, size=10, color=NAVY, bold=True)
    for label, value in data:
        cells = table.add_row().cells
        set_cell_shading(cells[0], LIGHT_GRAY)
        r = cells[0].paragraphs[0].add_run(label)
        set_run_font(r, size=10, color=NAVY, bold=True)
        r2 = cells[1].paragraphs[0].add_run(value if value else "\n\n")
        set_run_font(r2, size=10, color="777777", italic=bool(value == ""))
    set_table_geometry(table, [2400, 6960])
    doc.add_paragraph()


def build():
    doc = Document()
    style_doc(doc)
    add_header_footer(doc)

    add_title(doc, "Digital Forgetting Index", "Project Progress Evidence and Implementation Record")
    add_meta_table(doc)
    add_callout(doc, "Purpose", "This document records the current implementation of the Digital Forgetting Index model, explains how the code maps to the research paper, and leaves structured space for real execution results to be added later.")

    doc.add_heading("Executive summary", level=1)
    doc.add_paragraph("The Digital Forgetting Index (DFI) is a composite score for measuring how completely an AI system has forgotten data belonging to a subject after an erasure or unlearning request. The implementation evaluates four surfaces: datasets, databases, vector stores, and trained model parameters. Each surface is assessed with a layer-appropriate probe, converted to a 0-100 score, and combined using the paper's default 20/20/30/30 weighting scheme.")
    doc.add_paragraph("The current repository contains a working Python reference implementation, a deterministic synthetic benchmark, unit tests, documentation, and an audit-report object. Real production or experimental results have not yet been inserted; those results should be added in Section 12 after the model is run against the intended system.")

    doc.add_heading("1. Project objective and research basis", level=1)
    doc.add_paragraph("Modern AI pipelines can propagate personal data beyond the primary database. A deleted record may still remain in raw files, exports, replicas, caches, backups, embeddings, or model parameters. The research paper addresses this problem by replacing a binary deleted/not-deleted decision with a continuous and auditable forgetting score.")
    add_bullet(doc, "Measure residual personal information across multiple AI storage and model surfaces.")
    add_bullet(doc, "Use active probes rather than relying only on administrative deletion logs.")
    add_bullet(doc, "Report one composite DFI score while preserving layer-level diagnostic scores.")
    add_bullet(doc, "Make the measurement reproducible through explicit thresholds, weights, scope, timestamps, and probe versions.")
    add_callout(doc, "Interpretation boundary", "DFI is an empirical audit score. It is not proof of deletion, not equivalent to differential privacy, and not a legal compliance decision by itself.", fill="FFF8E8", color=GOLD)

    doc.add_heading("2. Repository structure", level=1)
    add_table(doc, ["Path", "Purpose", "Evidence status"], [
        ("dfi/core.py", "Implements formulas, weights, governance bands, validation, and audit reports.", "Implemented and tested"),
        ("dfi/probes.py", "Implements dataset, database, vector, and model probe interfaces.", "Implemented; production connectors pending"),
        ("dfi/synthetic.py", "Creates a deterministic benchmark across six unlearning strategies.", "Implemented"),
        ("dfi/cli.py", "Provides the benchmark command-line entry point.", "Implemented"),
        ("dfi/__init__.py", "Exports the primary public API.", "Implemented"),
        ("tests/test_dfi.py", "Checks formulas, fuzzy matching, vector detection, and validation.", "5 tests passing"),
        ("README.md", "Documents usage, formulas, assumptions, and limitations.", "Implemented"),
        ("pyproject.toml", "Defines the Python project metadata and CLI entry point.", "Implemented"),
    ], [2100, 5100, 2160])

    doc.add_heading("3. End-to-end workflow", level=1)
    doc.add_paragraph("The implementation follows this measurement sequence:")
    for text in [
        "Identify the erased subject and define the subject attributes used for probing.",
        "Establish the original baseline: records, stores, embeddings, and model exposure before deletion.",
        "Execute the deletion or unlearning strategy.",
        "Scan datasets and secondary copies for exact or near-duplicate records.",
        "Scan databases, replicas, logs, caches, and backups for residual records.",
        "Query the vector store using subject-derived embeddings and measure attributable top-k leakage.",
        "Run membership-inference and targeted extraction evaluation against the model.",
        "Convert raw signals into D-RS, DB-RS, V-RS, and M-RS.",
        "Apply the default weights and create the final DFI report.",
        "Review the result together with raw evidence, scope, thresholds, and limitations.",
    ]:
        add_number(doc, text)

    doc.add_heading("4. Mathematical model", level=1)
    add_table(doc, ["Layer", "Formula", "Meaning"], [
        ("Dataset", "D-RS = 100 x (1 - N_match / N_total)", "Residual exact or fuzzy matches across dataset copies."),
        ("Database", "DB-RS = 100 x (1 - R_res / R_orig)", "Residual records across database-related stores."),
        ("Vector store", "V-RS = 100 x (1 - I_top-k)", "Fraction of subject-attributable top-k retrieval leakage."),
        ("Model", "M-RS = 100 x (1 - 0.5 x (A_mia + A_ext))", "Normalized membership advantage plus extraction success."),
        ("Composite", "DFI = 0.20D-RS + 0.20DB-RS + 0.30V-RS + 0.30M-RS", "Weighted system-level forgetting score."),
    ], [1700, 4000, 3660])
    doc.add_paragraph("All final layer scores are constrained to the interval 0-100. A higher score indicates less detectable residue, not absolute absence of information.")

    doc.add_heading("5. Layer-by-layer implementation", level=1)
    doc.add_heading("5.1 Dataset residue", level=2)
    doc.add_paragraph("DatasetProbe accepts named copies such as raw files, CSV exports, archives, staging extracts, and backups. It canonicalizes selected fields by trimming whitespace and case-folding text, then applies a configurable fuzzy-match threshold. The residual count is summed across the declared copies and divided by the original baseline count.")
    doc.add_heading("5.2 Database residue", level=2)
    doc.add_paragraph("DatabaseProbe uses the same matching logic but treats operational stores as a separate audit scope. The caller may provide primary tables, replicas, audit logs, caches, and backups. This matters because deleting the primary row does not automatically prove that secondary stores are clean.")
    doc.add_heading("5.3 Vector-store residue", level=2)
    doc.add_paragraph("VectorProbe computes cosine similarity between subject-derived query embeddings and stored embeddings. It ranks the results, keeps the top-k items, and counts results that both exceed the similarity threshold and are attributable to the erased subject. The denominator is the number of query slots: number of queries multiplied by k.")
    doc.add_heading("5.4 Model residue", level=2)
    doc.add_paragraph("ModelProbe receives measured outputs from an external membership-inference or extraction harness. It converts raw MIA advantage into the paper's normalized advantage using A_mia = 2 x raw advantage, calculates the extraction success rate, and applies the model residue formula. The current reference code intentionally does not pretend that one generic attack works for every model architecture.")

    doc.add_heading("6. Report and audit metadata", level=1)
    doc.add_paragraph("compute_dfi() returns a DFIReport containing the four layer scores, composite score, governance band, weights, probe version, measurement timestamp, scope, raw signals, previous hash, and current SHA-256 report hash. The previous_hash field can be used to chain measurements over time.")
    add_table(doc, ["Metadata field", "Why it matters"], [
        ("Probe version", "Probe strength can change over time; scores from different versions may not be directly comparable."),
        ("Scope", "Documents which stores, copies, models, and retention windows were included or excluded."),
        ("Raw signals", "Allows an auditor to reconstruct how each sub-score was produced."),
        ("Measurement time", "DFI is point-in-time and can change after retraining or new ingestion."),
        ("Report hash", "Supports integrity checks and chained audit history."),
    ], [2200, 7160])

    doc.add_heading("7. How to run the current implementation", level=1)
    doc.add_paragraph("From the project directory, execute the following commands:")
    add_callout(doc, "Command 1", "python -m dfi.cli benchmark")
    add_callout(doc, "Command 2", "python -m unittest discover -v")
    doc.add_paragraph("The first command prints the synthetic benchmark as JSON. The second command runs the five automated tests. The expected validation status is all tests passing.")

    doc.add_heading("8. Synthetic benchmark evidence", level=1)
    doc.add_paragraph("The synthetic benchmark is a controlled demonstration rather than a measurement from a production system. It tests whether DFI increases as the assumed erasure strategy becomes stronger.")
    add_table(doc, ["Strategy", "Expected role", "Current illustrative DFI"], [
        ("No action", "Baseline with high residue", "3.00"),
        ("Database only", "Primary deletion without full unlearning", "43.75"),
        ("Approximate unlearning", "Partial model and storage reduction", "60.70"),
        ("SISA exact shard", "Retrain the affected shard", "79.15"),
        ("Certified removal", "Stronger model removal guarantee", "89.65"),
        ("SISA + vector purge", "Strongest combined synthetic strategy", "95.25"),
    ], [2600, 4800, 1960])
    add_callout(doc, "Evidence conclusion", "The synthetic sequence is monotonic: 3.00 < 43.75 < 60.70 < 79.15 < 89.65 < 95.25. This validates the intended direction of the scoring framework, but it does not validate real-world attack resistance.", fill="EAF4EA", color="1F5E2C")

    doc.add_heading("9. Worked calculation", level=1)
    doc.add_paragraph("The paper's worked example uses these layer scores:")
    add_table(doc, ["Layer", "Score", "Weight", "Contribution"], [
        ("Dataset", "97", "0.20", "19.40"),
        ("Database", "95", "0.20", "19.00"),
        ("Vector store", "88", "0.30", "26.40"),
        ("Model", "87", "0.30", "26.10"),
        ("Total", "-", "1.00", "90.90; paper rounds to 91.0"),
    ], [2200, 1400, 1500, 4260])
    doc.add_paragraph("This example shows why the report must retain sub-scores. A high composite can still hide a weaker model or vector score that requires engineering attention.")

    doc.add_heading("10. Validation completed so far", level=1)
    add_table(doc, ["Check", "Purpose", "Status"], [
        ("Formula test", "Reproduces the paper's worked example.", "Passed"),
        ("Model normalization test", "Checks raw MIA advantage and extraction conversion.", "Passed"),
        ("Dataset fuzzy-match test", "Detects near-duplicate residual records.", "Passed"),
        ("Vector top-k test", "Detects attributable high-similarity results.", "Passed"),
        ("Input validation test", "Rejects out-of-range scores.", "Passed"),
        ("Python compilation/import", "Checks source files load correctly.", "Passed"),
        ("Real-system execution", "Measures the target deployment.", "Pending"),
    ], [2600, 5000, 1760])

    doc.add_heading("11. Current limitations and required extensions", level=1)
    add_bullet(doc, "The reference implementation does not directly connect to production databases, vector databases, cloud backups, or model-serving endpoints.")
    add_bullet(doc, "The synthetic benchmark values are manually specified to test expected monotonicity; they are not observed measurements.")
    add_bullet(doc, "Real MIA and extraction probes must be chosen and calibrated for the model architecture and threat model.")
    add_bullet(doc, "Cross-layer interactions are not yet modeled. For example, a leaked database record might make a model extraction attack stronger.")
    add_bullet(doc, "The default weights are research defaults and may require domain-specific justification for healthcare, finance, or other high-risk applications.")
    add_bullet(doc, "The current probes are primarily designed for text and tabular data; multimodal extension remains future work.")

    doc.add_page_break()
    doc.add_heading("12. Results to be added after running the model", level=1)
    add_callout(doc, "Instructions", "Replace the blank fields below with the real measurement date, system versions, probe settings, raw signals, layer scores, final DFI, governance band, and interpretation.", fill="FFF8E8", color=GOLD)
    add_result_space(doc, "12.1 Execution record", rows=3)
    add_result_space(doc, "12.2 Layer-level results", rows=5)
    add_result_space(doc, "12.3 Composite result and interpretation", rows=4)
    add_result_space(doc, "12.4 Evidence attachments and reviewer notes", rows=4)

    doc.add_heading("13. Suggested final result narrative", level=1)
    doc.add_paragraph("After running the model, use the following structure to write the final conclusion:")
    for text in [
        "The measured DFI was [insert score] out of 100, corresponding to the [insert governance band] band.",
        "The strongest layer was [insert layer] with a score of [insert score], indicating [insert interpretation].",
        "The weakest layer was [insert layer] with a score of [insert score], indicating that [insert residual risk] remains.",
        "The main evidence was [insert counts, retrieval leakage, MIA advantage, or extraction results].",
        "The result should be interpreted within the documented scope and probe version; it is not proof of complete deletion.",
    ]:
        add_bullet(doc, text)

    doc.add_heading("14. Source and evidence references", level=1)
    doc.add_paragraph("Primary research source: Draft 1.pdf, Digital Forgetting Index: A Quantifiable Framework for Measuring AI Data Erasure.")
    doc.add_paragraph("Implementation source files: the Python modules and tests contained in this project repository. The code is intended to make the paper's framework executable and auditable while keeping real-system integrations explicit.")

    doc.core_properties.title = "Digital Forgetting Index - Project Progress Evidence"
    doc.core_properties.subject = "Implementation progress, workflow, formulas, validation, and pending results"
    doc.core_properties.author = "DFI Project"
    doc.core_properties.comments = "Generated project evidence document"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
